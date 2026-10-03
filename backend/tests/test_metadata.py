import io
from hashlib import sha256
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from docx import Document
from pypdf import PdfWriter

from app.auth.hashing import get_password_hash
from app.main import app
from app.models.case import Case
from app.models.user import User
from app.services.metadata_service import MetadataService
from tests.conftest import TestSessionLocal

client = TestClient(app)


@pytest.fixture
def temp_files(tmp_path):
    jpg_path = tmp_path / "sample.jpg"
    png_path = tmp_path / "sample.png"
    pdf_path = tmp_path / "sample.pdf"
    docx_path = tmp_path / "sample.docx"
    txt_path = tmp_path / "notes.txt"
    corrupted = tmp_path / "corrupted.pdf"
    corrupted.write_bytes(b"not a valid pdf")

    img = Image.new("RGB", (120, 80), color="blue")
    img.save(jpg_path)
    img.save(png_path)

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf_path.open("wb") as handle:
        writer.write(handle)

    doc = Document()
    doc.add_paragraph("Evidence metadata test")
    doc.core_properties.title = "Case Report"
    doc.core_properties.subject = "Evidence"
    doc.core_properties.author = "Analyst"
    doc.core_properties.keywords = "forensic, validation"
    doc.core_properties.comments = "Sample DOCX metadata"
    doc.core_properties.last_modified_by = "System"
    doc.save(docx_path)

    txt_path.write_text("plain text evidence", encoding="utf-8")

    return {
        "jpg": jpg_path,
        "png": png_path,
        "pdf": pdf_path,
        "docx": docx_path,
        "txt": txt_path,
        "corrupted": corrupted,
    }


def test_jpeg_metadata_extraction(temp_files):
    result = MetadataService.extract_metadata_for_file(temp_files["jpg"], filename="sample.jpg")
    assert result["file_type"] in {"image/jpeg", "image"}
    assert result["extraction_status"] in {"success", "partial"}
    assert "metadata" in result
    assert result["metadata"].get("format") in {"JPEG", "JPEG"}


def test_png_metadata_extraction(temp_files):
    result = MetadataService.extract_metadata_for_file(temp_files["png"], filename="sample.png")
    assert result["file_type"] in {"image/png", "image"}
    assert result["extraction_status"] in {"success", "partial"}
    assert result["metadata"].get("width") is not None
    assert result["metadata"].get("height") is not None


def test_pdf_metadata_extraction(temp_files):
    result = MetadataService.extract_metadata_for_file(temp_files["pdf"], filename="sample.pdf")
    assert result["file_type"] == "application/pdf"
    assert result["extraction_status"] in {"success", "partial"}
    assert isinstance(result["metadata"], dict)


def test_docx_metadata_extraction(temp_files):
    result = MetadataService.extract_metadata_for_file(temp_files["docx"], filename="sample.docx")
    assert result["file_type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert result["extraction_status"] in {"success", "partial"}
    assert result["metadata"].get("author") in {"Analyst", "System"}


def test_docx_metadata_endpoint_returns_supported_response():
    db = TestSessionLocal()
    user = db.query(User).filter(User.email == "docx-metadata-user@example.com").first()
    if user is None:
        user = User(
            username="docx-metadata-user",
            email="docx-metadata-user@example.com",
            hashed_password=get_password_hash("StrongPass123!"),
            is_active=True,
        )
        db.add(user)
        db.commit()
        db.refresh(user)

    case = db.query(Case).filter(Case.case_number == "CASE-METADATA-DOCX").first()
    if case is None:
        case = Case(
            case_number="CASE-METADATA-DOCX",
            title="DOCX Metadata Case",
            description="DOCX metadata route verification",
            created_by=user.id,
            status="open",
        )
        db.add(case)
        db.commit()
        db.refresh(case)

    file_path = Path("uploads") / f"case_{case.id}" / "metadata_case.docx"
    file_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_paragraph("Case evidence report")
    doc.core_properties.title = "Case Report"
    doc.core_properties.subject = "Evidence"
    doc.core_properties.author = "Analyst"
    doc.core_properties.keywords = "forensic, validation"
    doc.core_properties.comments = "Sample DOCX metadata"
    doc.core_properties.last_modified_by = "System"
    doc.save(file_path)

    digest = sha256(file_path.read_bytes()).hexdigest()

    from app.models.evidence import Evidence

    evidence = Evidence(
        case_id=case.id,
        owner_id=user.id,
        file_name="metadata_case.docx",
        file_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_size=file_path.stat().st_size,
        storage_path=str(file_path),
        sha256_hash=digest,
        description="docx evidence",
    )
    db.add(evidence)
    db.commit()
    db.refresh(evidence)
    before_hash = evidence.sha256_hash
    db.close()

    login_response = client.post(
        "/api/auth/login",
        json={"username": "docx-metadata-user", "password": "StrongPass123!"},
    )
    assert login_response.status_code == 200, login_response.text
    token = login_response.json()["access_token"]

    response = client.get(
        f"/api/metadata/evidence/{evidence.id}",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["file_type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert payload["mime_type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert payload["extraction_status"] in {"success", "partial"}
    assert payload["metadata"]["filename"] == "metadata_case.docx"
    assert payload["metadata"]["file_size"] == file_path.stat().st_size
    assert payload["metadata"]["author"] == "Analyst"
    assert payload["metadata"]["title"] == "Case Report"
    assert payload["extraction_errors"] == []

    db = TestSessionLocal()
    stored = db.query(Evidence).filter(Evidence.id == evidence.id).first()
    assert stored is not None
    assert stored.sha256_hash == before_hash
    db.close()


def test_text_metadata_extraction(temp_files):
    result = MetadataService.extract_metadata_for_file(temp_files["txt"], filename="notes.txt", mime_type="text/plain")
    assert result["file_type"] == "text/plain"
    assert result["mime_type"] == "text/plain"
    assert result["extraction_status"] in {"success", "partial"}
    metadata = result["metadata"]
    assert metadata["filename"] == "notes.txt"
    assert metadata["extension"] == ".txt"
    assert metadata["mime_type"] == "text/plain"
    assert metadata["file_size"] == temp_files["txt"].stat().st_size
    assert metadata["character_count"] == len("plain text evidence")
    assert metadata["line_count"] == 1
    assert metadata["word_count"] == 3
    assert "encoding" in metadata
    assert metadata["created_at"] is not None or metadata.get("created") is not None
    assert metadata["modified_at"] is not None or metadata.get("modified") is not None


def test_text_metadata_endpoint_returns_supported_response():
    db = TestSessionLocal()
    user = db.query(User).filter(User.email == "metadata-user@example.com").first()
    if user is None:
        user = User(
            username="metadata-user",
            email="metadata-user@example.com",
            hashed_password=get_password_hash("StrongPass123!"),
            is_active=True,
        )
        db.add(user)
        db.commit()
        db.refresh(user)

    case = db.query(Case).filter(Case.case_number == "CASE-METADATA-TXT").first()
    if case is None:
        case = Case(
            case_number="CASE-METADATA-TXT",
            title="Text Metadata Case",
            description="Text metadata route verification",
            created_by=user.id,
            status="open",
        )
        db.add(case)
        db.commit()
        db.refresh(case)

    file_path = Path("uploads") / f"case_{case.id}" / "metadata_case.txt"
    file_path.parent.mkdir(parents=True, exist_ok=True)
    content = "alpha beta gamma\nline two\n"
    file_path.write_text(content, encoding="utf-8")
    digest = sha256(file_path.read_bytes()).hexdigest()

    from app.models.evidence import Evidence

    evidence = Evidence(
        case_id=case.id,
        owner_id=user.id,
        file_name="metadata_case.txt",
        file_type="text/plain",
        file_size=file_path.stat().st_size,
        storage_path=str(file_path),
        sha256_hash=digest,
        description="text evidence",
    )
    db.add(evidence)
    db.commit()
    db.refresh(evidence)
    db.close()

    login_response = client.post(
        "/api/auth/login",
        json={"username": "metadata-user", "password": "StrongPass123!"},
    )
    assert login_response.status_code == 200, login_response.text
    token = login_response.json()["access_token"]

    response = client.get(
        f"/api/metadata/evidence/{evidence.id}",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["file_type"] == "text/plain"
    assert payload["mime_type"] == "text/plain"
    assert payload["extraction_status"] in {"success", "partial"}
    assert payload["metadata"]["filename"] == "metadata_case.txt"
    assert payload["metadata"]["word_count"] == 5
    assert payload["metadata"]["line_count"] == 2


def test_pdf_metadata_endpoint_returns_supported_response():
    db = TestSessionLocal()
    user = db.query(User).filter(User.email == "pdf-metadata-user@example.com").first()
    if user is None:
        user = User(
            username="pdf-metadata-user",
            email="pdf-metadata-user@example.com",
            hashed_password=get_password_hash("StrongPass123!"),
            is_active=True,
        )
        db.add(user)
        db.commit()
        db.refresh(user)

    case = db.query(Case).filter(Case.case_number == "CASE-METADATA-PDF").first()
    if case is None:
        case = Case(
            case_number="CASE-METADATA-PDF",
            title="PDF Metadata Case",
            description="PDF metadata route verification",
            created_by=user.id,
            status="open",
        )
        db.add(case)
        db.commit()
        db.refresh(case)

    file_path = Path("uploads") / f"case_{case.id}" / "metadata_case.pdf"
    file_path.parent.mkdir(parents=True, exist_ok=True)

    writer = PdfWriter()
    page = writer.add_blank_page(width=72, height=72)
    writer.add_metadata({
        "/Title": "Case Report",
        "/Author": "Analyst",
        "/Subject": "Evidence",
        "/Producer": "Digital Evidence Integrity",
        "/Creator": "pytest",
        "/CreationDate": "D:20240101010101Z",
        "/ModDate": "D:20240102020202Z",
    })
    with file_path.open("wb") as handle:
        writer.write(handle)

    digest = sha256(file_path.read_bytes()).hexdigest()

    from app.models.evidence import Evidence

    evidence = Evidence(
        case_id=case.id,
        owner_id=user.id,
        file_name="metadata_case.pdf",
        file_type="application/pdf",
        file_size=file_path.stat().st_size,
        storage_path=str(file_path),
        sha256_hash=digest,
        description="pdf evidence",
    )
    db.add(evidence)
    db.commit()
    db.refresh(evidence)
    db.close()

    login_response = client.post(
        "/api/auth/login",
        json={"username": "pdf-metadata-user", "password": "StrongPass123!"},
    )
    assert login_response.status_code == 200, login_response.text
    token = login_response.json()["access_token"]

    response = client.get(
        f"/api/metadata/evidence/{evidence.id}",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["file_type"] == "application/pdf"
    assert payload["mime_type"] == "application/pdf"
    assert payload["extraction_status"] in {"success", "partial"}
    assert payload["metadata"]["filename"] == "metadata_case.pdf"
    assert payload["metadata"]["file_size"] == file_path.stat().st_size
    assert payload["metadata"]["title"] == "Case Report"
    assert payload["metadata"]["author"] == "Analyst"
    assert payload["metadata"]["page_count"] == 1


def test_missing_file():
    result = MetadataService.extract_metadata_for_file("/tmp/definitely_missing.file", filename="missing.file")
    assert result["extraction_status"] == "failed"
    assert result["extraction_errors"]


def test_corrupted_file_handling(temp_files):
    result = MetadataService.extract_metadata_for_file(temp_files["corrupted"], filename="corrupted.pdf")
    assert result["extraction_status"] in {"failed", "partial"}
    assert isinstance(result["extraction_errors"], list)


def test_metadata_response_serialization(temp_files):
    result = MetadataService.extract_metadata_for_file(temp_files["jpg"], filename="sample.jpg")
    assert isinstance(result["metadata"], dict)
    assert isinstance(result["extraction_errors"], list)
    assert isinstance(result["extraction_status"], str)
